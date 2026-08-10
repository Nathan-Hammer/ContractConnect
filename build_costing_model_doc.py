from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK

OUT = r"C:\Users\n.hammerslagt\Downloads\ContractConnect\ContractConnect_Costing_Model.docx"

NAVY = "123B3A"
TEAL = "1D6B66"
GOLD = "D9A441"
LIGHT = "EEF5F4"
PALE = "F5F7F7"
MID = "D8E4E2"
TEXT = "253331"
MUTED = "64706E"
WHITE = "FFFFFF"

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.8)
sec.bottom_margin = Inches(0.75)
sec.left_margin = Inches(0.85)
sec.right_margin = Inches(0.85)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for name, size, color, before, after in [
    ("Title", 28, NAVY, 0, 8),
    ("Subtitle", 12.5, MUTED, 0, 12),
    ("Heading 1", 16, NAVY, 16, 8),
    ("Heading 2", 12.5, TEAL, 11, 5),
    ("Heading 3", 11, NAVY, 8, 4),
]:
    st = styles[name]
    st.font.name = "Aptos Display" if name in ("Title", "Heading 1") else "Aptos"
    st.font.size = Pt(size)
    st.font.bold = name != "Subtitle"
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for list_name in ("List Bullet", "List Number"):
    st = styles[list_name]
    st.font.name = "Aptos"
    st.font.size = Pt(10.5)
    st.paragraph_format.left_indent = Inches(0.5)
    st.paragraph_format.first_line_indent = Inches(-0.25)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.10

if "Small Note" not in styles:
    note = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
else:
    note = styles["Small Note"]
note.font.name = "Aptos"
note.font.size = Pt(8.5)
note.font.color.rgb = RGBColor.from_string(MUTED)
note.paragraph_format.space_after = Pt(4)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_table_borders(table, color="CDD9D7", size="6"):
    tblPr = table._tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def set_col_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = Inches(width)
            tcPr = row.cells[i]._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            tcW.set(qn("w:w"), str(round(width * 1440)))
            tcW.set(qn("w:type"), "dxa")
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    tblW.set(qn("w:w"), str(round(sum(widths) * 1440)))
    tblW.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(round(width * 1440)))
        grid.append(col)

def add_table(headers, rows, widths, aligns=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_col_widths(table, widths)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.runs[0]
        r.bold = True
        r.font.name = "Aptos"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string(WHITE)
    set_repeat_table_header(table.rows[0])
    for ri, row_data in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].text = str(value)
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2 == 1:
                set_cell_shading(cells[i], PALE)
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if aligns and aligns[i]:
                p.alignment = aligns[i]
            for r in p.runs:
                r.font.name = "Aptos"
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor.from_string(TEXT)
    set_table_borders(table)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table

def add_field(paragraph, field_code):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = field_code
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)

def add_callout(label, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_col_widths(table, [6.8])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    set_cell_margins(cell, top=150, bottom=150, start=180, end=180)
    set_table_borders(table, color=MID, size="8")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + " ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)
    r2 = p.add_run(text)
    r2.font.color.rgb = RGBColor.from_string(TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

header = sec.header
hp = header.paragraphs[0]
hp.text = "CONTRACTCONNECT  |  COMMERCIAL PRICING MODEL"
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in hp.runs:
    r.font.name = "Aptos"
    r.font.size = Pt(8)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)

footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("ContractConnect - Confidential  |  Page ")
fr.font.name = "Aptos"
fr.font.size = Pt(8)
fr.font.color.rgb = RGBColor.from_string(MUTED)
add_field(fp, "PAGE")

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(70)
p.paragraph_format.space_after = Pt(8)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("CONTRACTCONNECT")
r.font.name = "Aptos"
r.font.size = Pt(12)
r.font.bold = True
r.font.color.rgb = RGBColor.from_string(GOLD)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Costing & Commercial Pricing Model")

p = doc.add_paragraph(style="Subtitle")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("A practical go-to-market model for a contract relationship management solution")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(18)
p.paragraph_format.space_after = Pt(26)
r = p.add_run("Prepared for product presentation and commercial planning\n6 August 2026")
r.font.name = "Aptos"
r.font.size = Pt(10)
r.font.color.rgb = RGBColor.from_string(MUTED)

add_callout("Recommended model:", "A fixed monthly platform fee with capacity bands for users, active contracts and storage. This gives customers predictable costs while protecting ContractConnect's margins as usage grows.")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(54)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("All prices are stated in Namibian dollars and exclude VAT.")
r.italic = True
r.font.size = Pt(9)
r.font.color.rgb = RGBColor.from_string(MUTED)

doc.add_page_break()

doc.add_heading("1. Executive recommendation", level=1)
doc.add_paragraph("ContractConnect should be positioned as an affordable, locally supported contract relationship management platform for small and medium-sized organisations. The recommended commercial model combines the predictability of a platform subscription with sensible capacity limits.")

for text in [
    "Use contract and storage capacity as the main value metric, with included user allowances.",
    "Lead with the Professional plan as the primary commercial offer.",
    "Charge implementation separately so onboarding and data preparation are properly funded.",
    "Use a paid pilot to reduce customer risk without permanently discounting the product.",
    "Initially deploy a dedicated customer environment until full multi-tenant isolation is implemented.",
]:
    doc.add_paragraph(text, style="List Bullet")

doc.add_heading("2. Subscription pricing", level=1)
add_table(
    ["Plan", "Monthly", "Users", "Active contracts", "Storage", "Best suited for"],
    [
        ["Essential", "N$3,500", "5", "100", "5 GB", "Small businesses"],
        ["Professional", "N$6,900", "20", "500", "25 GB", "Growing teams"],
        ["Business", "N$12,500", "50", "2,000", "100 GB", "Larger organisations"],
        ["Enterprise", "From N$22,000", "Custom", "Custom", "Custom", "Regulated or complex organisations"],
    ],
    [1.05, 1.0, 0.55, 1.05, 0.7, 2.45],
    [None, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, None],
)

doc.add_heading("Included in all plans", level=2)
for text in [
    "Company and contact management; contract records and lifecycle tracking",
    "Interaction history, follow-ups and renewal reminders",
    "Dashboard, reports, search, filters and exports",
    "Authentication, user profiles and document management",
    "Standard customer support",
]:
    doc.add_paragraph(text, style="List Bullet")

doc.add_heading("Plan differentiation", level=2)
doc.add_paragraph("Professional and higher plans may include advanced roles, private documents, audit history, saved views, enhanced reporting, priority support and additional training. Enterprise pricing should cover dedicated hosting, customised security controls, service-level commitments and bespoke development.")

doc.add_heading("3. Once-off implementation fees", level=1)
add_table(
    ["Plan", "Implementation fee"],
    [["Essential", "N$12,000"], ["Professional", "N$25,000"], ["Business", "N$55,000"], ["Enterprise", "From N$100,000"]],
    [3.4, 3.4],
    [None, WD_ALIGN_PARAGRAPH.RIGHT],
)
doc.add_paragraph("The implementation fee covers customer-environment setup, branding and configuration, roles, data-import templates, initial approved-data import, administrator training and go-live support.")
add_callout("Scope control:", "Significant data cleaning, custom reports, integrations and workflow changes should be quoted separately.")

doc.add_heading("4. Optional services", level=1)
add_table(
    ["Service", "Suggested charge"],
    [
        ["Additional 5-user block", "N$600/month"],
        ["Additional 10 GB storage", "N$300/month"],
        ["Additional training session", "N$2,500"],
        ["Data migration", "N$950/hour or fixed quotation"],
        ["Custom development", "N$950-N$1,200/hour"],
        ["Full consulting day", "N$7,500-N$9,000"],
        ["Premium support", "From N$2,500/month"],
        ["Custom-branded deployment", "From N$15,000 once-off"],
    ],
    [3.7, 3.1],
    [None, WD_ALIGN_PARAGRAPH.RIGHT],
)
doc.add_paragraph("Calendar, email, electronic-signature and AI integrations should remain separately priced future add-ons until each capability is production-ready.")

doc.add_heading("5. Pilot offer", level=1)
add_callout("60-day paid pilot - N$15,000:", "Up to 10 users and 100 contracts. The pilot fee is credited against implementation if the customer signs an annual subscription.")
doc.add_paragraph("The pilot gives the customer a controlled evaluation while preserving the product's commercial value and funding onboarding support.")

doc.add_heading("6. Annual contract terms", level=1)
for text in [
    "Monthly payment: standard published price.",
    "Annual prepayment: 15% subscription discount.",
    "Minimum initial term: 12 months following any pilot.",
    "Annual increase: Namibian CPI or 7%, whichever is lower.",
    "Migration, customisation and third-party service fees are excluded from subscription pricing.",
]:
    doc.add_paragraph(text, style="List Bullet")

doc.add_heading("Professional plan - first-year example", level=2)
add_table(
    ["Cost component", "Amount"],
    [
        ["Subscription: N$6,900 x 12", "N$82,800"],
        ["Annual prepayment discount (15%)", "(N$12,420)"],
        ["Discounted annual subscription", "N$70,380"],
        ["Implementation", "N$25,000"],
        ["First-year total", "N$95,380"],
        ["Subsequent annual subscription", "N$70,380"],
    ],
    [4.6, 2.2],
    [None, WD_ALIGN_PARAGRAPH.RIGHT],
)

doc.add_heading("7. Unit economics and pricing floor", level=1)
doc.add_paragraph("Pricing should target a gross margin of approximately 60-70% after direct hosting, support and maintenance costs. The commercial pricing floor can be calculated as:")
add_callout("Pricing formula:", "Minimum monthly price = direct monthly customer cost / (1 - target gross margin)")
doc.add_paragraph("For example, direct monthly customer costs of N$1,500 at a 70% target gross margin require a minimum selling price of approximately N$5,000 per month. The Essential tier may operate at a lower initial margin, while Professional and Business protect the blended portfolio margin.")

doc.add_heading("Suggested internal cost allocation", level=2)
add_table(
    ["Cost category", "Planning allowance"],
    [
        ["Hosting, database, backups and monitoring", "N$900-N$1,300/customer/month"],
        ["Routine support allowance", "N$400-N$1,000/customer/month"],
        ["Maintenance and security reserve", "10-15% of subscription revenue"],
        ["Billing, administration and payment costs", "3-5% of revenue"],
        ["Sales commission", "10% of first-year subscription or one month MRR"],
    ],
    [4.15, 2.65],
    [None, WD_ALIGN_PARAGRAPH.RIGHT],
)

doc.add_heading("8. Market positioning", level=1)
doc.add_paragraph("ContractConnect combines CRM-style relationship management with contract lifecycle tracking. International CRM platforms commonly price per user, while dedicated contract-management products often price by contract volume. A capacity-based platform subscription therefore provides an appropriate middle ground for the solution.")

add_table(
    ["Reference product", "Observed pricing approach"],
    [
        ["Pipedrive", "Per-user CRM subscriptions; entry pricing from about US$14/user/month annually"],
        ["HubSpot Sales Hub", "Seat-based CRM pricing across Starter, Professional and Enterprise editions"],
        ["ContractSafe", "Contract-volume pricing with unlimited users; entry pricing around US$450/month annually"],
        ["Supabase", "Backend platform; Pro plan starts at US$25/month before usage charges"],
        ["Vercel", "Application hosting; Pro plan starts at US$20/month before usage charges"],
    ],
    [1.7, 5.1],
)

doc.add_heading("9. Technical commercial-readiness note", level=1)
add_callout("Important:", "The current ContractConnect environment should initially be sold as a dedicated deployment for each customer. Before unrelated organisations share one database, full tenant isolation must be implemented and tested so one customer's users can never access another customer's records.")
doc.add_paragraph("This deployment approach should be reflected in hosting estimates, support commitments and implementation pricing. Once verified multi-tenancy is available, lower-cost shared infrastructure can be introduced for Essential and Professional customers.")

doc.add_heading("10. Recommended launch offer", level=1)
doc.add_paragraph("The Professional plan should be the flagship launch package. It provides a credible feature and capacity allowance for operational teams while producing sufficient recurring revenue to fund dedicated infrastructure, support and continued product development.")
for text in [
    "Flagship subscription: Professional at N$6,900 per month.",
    "Implementation: N$25,000 once-off.",
    "Annual prepayment: N$70,380, excluding VAT.",
    "First-year total: N$95,380, excluding VAT.",
    "Entry path: 60-day pilot at N$15,000, credited on conversion.",
]:
    doc.add_paragraph(text, style="List Bullet")

doc.add_heading("Sources", level=1)
sources = [
    ("Pipedrive pricing", "https://www.pipedrive.com/en/pricing/professional-crm"),
    ("HubSpot Sales pricing", "https://www.hubspot.com/pricing/sales"),
    ("ContractSafe pricing", "https://www.contractsafe.com/pricing"),
    ("Supabase pricing", "https://supabase.com/pricing"),
    ("Vercel pricing", "https://vercel.com/pricing"),
]
for name, url in sources:
    p = doc.add_paragraph(style="Small Note")
    p.add_run(f"{name}: ").bold = True
    p.add_run(url)

doc.core_properties.title = "ContractConnect Costing and Commercial Pricing Model"
doc.core_properties.subject = "Proposed subscription, implementation and commercial pricing model"
doc.core_properties.author = "ContractConnect"
doc.core_properties.keywords = "ContractConnect, costing, pricing, SaaS, CRM, contract management"
doc.save(OUT)
print(OUT)
